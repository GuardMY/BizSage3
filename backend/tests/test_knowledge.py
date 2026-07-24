"""Lifecycle tests for versioned knowledge ingestion and report citations."""

import io
from datetime import datetime
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import (
    Base,
    KnowledgeAuditEvent,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeDocumentVersion,
    KnowledgeIngestionJob,
    KnowledgeRetrievalConfiguration,
    ReportEvidence,
)
from app.repository import SessionRepository
from app.services.knowledge import (
    EmbeddingService,
    EvidenceContext,
    KnowledgeIngestionService,
    KnowledgeRetrievalService,
    chunk_blocks,
    evidence_from_state,
    evidence_to_state,
    parse_document,
    persist_report_evidences,
    validate_report_citations,
)


class FakeStorage:
    def __init__(self, content: bytes):
        self.content = content

    async def get(self, storage_key: str) -> bytes:
        return self.content


class FakeEmbedder:
    async def embed(self, texts: list[str]) -> list[list[float]]:
        return [[0.1, 0.2, 0.3] for _ in texts]


class RecordingEmbedder(FakeEmbedder):
    def __init__(self):
        self.inputs: list[str] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.inputs.extend(texts)
        return await super().embed(texts)


class RecordingEmbeddingClient:
    def __init__(self):
        self.calls: list[list[str]] = []
        self.embeddings = self

    async def create(self, *, model, input, dimensions):
        self.calls.append(list(input))
        return SimpleNamespace(
            data=[
                SimpleNamespace(embedding=[float(len(self.calls)), float(index)])
                for index, _ in enumerate(input)
            ]
        )


class FakeVectorIndex:
    def __init__(self):
        self.indexed: list[tuple[str, int]] = []

    async def upsert_chunks(self, chunks, vectors, version) -> None:
        assert len(chunks) == len(vectors)
        self.indexed.append((version.id, len(chunks)))


class EmptyVectorIndex:
    async def search(self, vector, scene, limit):
        return []


class RecalledVectorIndex:
    def __init__(self, recalled: list[tuple[str, float]]):
        self.recalled = recalled

    async def search(self, vector, scene, limit):
        return self.recalled[:limit]


class RecordingVectorIndex(RecalledVectorIndex):
    def __init__(self, recalled: list[tuple[str, float]]):
        super().__init__(recalled)
        self.scenes: list[dict[str, str]] = []

    async def search(self, vector, scene, limit):
        self.scenes.append(dict(scene))
        return await super().search(vector, scene, limit)


def test_evidence_state_is_json_checkpoint_safe():
    evidence = EvidenceContext(
        chunk_id="chunk", version_id="version", document_id="document",
        document_title="餐饮方法", source_type="methodology", version_no=1,
        quote="原文片段", locator={"line_start": 3}, query="外卖转化", rank=1,
        retrieved_at=datetime(2026, 7, 16, 10, 0, 0),
    )

    restored = evidence_from_state([evidence_to_state(evidence)])

    assert restored == [evidence]


@pytest.fixture
async def knowledge_db_factory():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    yield factory
    await engine.dispose()


@pytest.mark.asyncio
async def test_ingestion_creates_reviewable_chunks(knowledge_db_factory):
    async with knowledge_db_factory() as db:
        document = KnowledgeDocument(title="外卖转化方法", status="draft")
        db.add(document)
        await db.flush()
        version = KnowledgeDocumentVersion(
            document_id=document.id,
            version_no=1,
            original_filename="delivery.md",
            content_type="text/markdown",
            source_type="methodology",
            sha256="a" * 64,
            storage_key="documents/version/delivery.md",
            status="draft",
            industry_tags=["餐饮"],
            sub_industry_tags=[],
            business_mode_tags=["外卖"],
            operating_stage_tags=["增长"],
        )
        db.add(version)
        await db.flush()
        db.add(KnowledgeIngestionJob(version_id=version.id, state="queued"))
        await db.commit()
        version_id = version.id

    vectors = FakeVectorIndex()
    service = KnowledgeIngestionService(
        storage=FakeStorage("# 转化诊断\n\n菜单曝光下降时，先检查曝光到下单的漏斗。".encode()),
        vector_index=vectors,
        embedding_service=FakeEmbedder(),
        session_factory=knowledge_db_factory,
    )
    async with knowledge_db_factory() as db:
        job_id = (await db.execute(
            select(KnowledgeIngestionJob.id).where(KnowledgeIngestionJob.version_id == version_id)
        )).scalar_one()
    await service.run(job_id, worker_id="test-worker")

    async with knowledge_db_factory() as db:
        version = await db.get(KnowledgeDocumentVersion, version_id)
        chunks = list((await db.execute(
            select(KnowledgeChunk).where(KnowledgeChunk.version_id == version_id)
        )).scalars().all())
        job = (await db.execute(
            select(KnowledgeIngestionJob).where(KnowledgeIngestionJob.version_id == version_id)
        )).scalar_one()

    assert version.status == "pending_review"
    assert job.state == "completed"
    assert job.worker_id == "test-worker"
    assert chunks and "曝光到下单" in chunks[0].content
    assert vectors.indexed == [(version_id, len(chunks))]


@pytest.mark.asyncio
async def test_embedding_service_batches_requests_to_provider(monkeypatch):
    monkeypatch.setattr("app.services.knowledge.embedding_api_key", lambda: "test-key")
    client = RecordingEmbeddingClient()
    service = EmbeddingService.__new__(EmbeddingService)
    service._client = client

    texts = [f"chunk-{index}" for index in range(12)]
    vectors = await service.embed(texts)

    assert [len(call) for call in client.calls] == [10, 2]
    assert len(vectors) == 12


def test_markdown_frontmatter_and_table_locators_are_precise():
    parsed = parse_document(
        "playbook.md",
        (
            b"---\n"
            b"title: Sample\n"
            b"---\n"
            b"# Overview\n\n"
            b"Intro line.\n\n"
            b"## Metrics\n\n"
            b"| KPI | Value |\n"
            b"| --- | --- |\n"
            b"| Conversion | 25% |\n"
            b"| AOV | 88 |\n"
        ),
    )

    chunks = chunk_blocks(parsed.blocks)

    assert all("title: Sample" not in text for text, _ in chunks)
    assert [locator["kind"] for _, locator in chunks] == ["paragraph", "table"]
    assert chunks[0][1]["heading_path"] == ["Overview"]
    assert chunks[0][1]["line_start"] == 6
    assert chunks[0][1]["line_end"] == 6
    table_locator = chunks[1][1]
    assert table_locator["heading_path"] == ["Overview", "Metrics"]
    assert table_locator["line_start"] == 10
    assert table_locator["line_end"] == 13
    assert table_locator["table_no"] == 1
    assert table_locator["row_start"] == 1
    assert table_locator["row_end"] == 2


def test_chunk_blocks_do_not_cross_heading_boundaries():
    parsed = parse_document(
        "sections.md",
        (
            b"# Guide\n\n"
            b"## Diagnose\n\n"
            b"Alpha section.\n\n"
            b"## Improve\n\n"
            b"Beta section.\n"
        ),
    )

    chunks = chunk_blocks(parsed.blocks)

    assert len(chunks) == 2
    assert chunks[0][1]["heading_path"] == ["Guide", "Diagnose"]
    assert chunks[1][1]["heading_path"] == ["Guide", "Improve"]
    assert "Beta section." not in chunks[0][0]
    assert "Alpha section." not in chunks[1][0]


def test_docx_parser_keeps_heading_paths_and_table_rows():
    document = Document()
    document.add_heading("Playbook", level=1)
    document.add_paragraph("Lead with diagnosis")
    document.add_heading("Checklist", level=2)
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Conversion"
    table.rows[1].cells[1].text = "25%"
    table.rows[2].cells[0].text = "AOV"
    table.rows[2].cells[1].text = "88"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document("playbook.docx", buffer.getvalue())

    assert [block.locator["kind"] for block in parsed.blocks] == ["paragraph", "table"]
    paragraph = parsed.blocks[0]
    assert paragraph.locator["heading_path"] == ["Playbook"]
    assert paragraph.locator["paragraph_start"] == 2
    assert paragraph.locator["paragraph_end"] == 2
    table_block = parsed.blocks[1]
    assert table_block.locator["heading_path"] == ["Playbook", "Checklist"]
    assert table_block.locator["table_no"] == 1
    assert table_block.locator["row_start"] == 1
    assert table_block.locator["row_end"] == 2

    chunks = chunk_blocks(parsed.blocks)

    assert len(chunks) == 2
    assert chunks[1][1]["table_no"] == 1
    assert chunks[1][1]["row_start"] == 1
    assert chunks[1][1]["row_end"] == 2


@pytest.mark.asyncio
async def test_ingestion_embeddings_include_document_title_and_headings(knowledge_db_factory):
    async with knowledge_db_factory() as db:
        document = KnowledgeDocument(title="Delivery Playbook", status="draft")
        db.add(document)
        await db.flush()
        version = KnowledgeDocumentVersion(
            document_id=document.id,
            version_no=1,
            original_filename="delivery.md",
            content_type="text/markdown",
            source_type="methodology",
            sha256="f" * 64,
            storage_key="documents/version/delivery.md",
            status="draft",
        )
        db.add(version)
        await db.flush()
        db.add(KnowledgeIngestionJob(version_id=version.id, state="queued"))
        await db.commit()
        version_id = version.id

    embedder = RecordingEmbedder()
    service = KnowledgeIngestionService(
        storage=FakeStorage(b"# Diagnosis\n\nCheck funnel health.\n"),
        vector_index=FakeVectorIndex(),
        embedding_service=embedder,
        session_factory=knowledge_db_factory,
    )
    async with knowledge_db_factory() as db:
        job_id = (await db.execute(
            select(KnowledgeIngestionJob.id).where(KnowledgeIngestionJob.version_id == version_id)
        )).scalar_one()
    await service.run(job_id, worker_id="test-worker")

    assert embedder.inputs == ["Delivery Playbook\nDiagnosis\nCheck funnel health."]


@pytest.mark.asyncio
async def test_retrieval_does_not_write_audit_event(knowledge_db_factory):
    service = KnowledgeRetrievalService(
        vector_index=EmptyVectorIndex(),
        embedding_service=FakeEmbedder(),
        session_factory=knowledge_db_factory,
    )
    service._ready = True

    assert await service.retrieve("机械制造业关键指标", {"industry": "机械制造业"}) == []

    async with knowledge_db_factory() as db:
        events = list((await db.execute(select(KnowledgeAuditEvent))).scalars().all())

    assert events == []


@pytest.mark.asyncio
async def test_retrieval_exposes_independent_scores_and_final_rank(knowledge_db_factory):
    async with knowledge_db_factory() as db:
        methodology = KnowledgeDocument(title="转化方法", status="published")
        case = KnowledgeDocument(title="经营案例", status="published")
        db.add_all([methodology, case])
        await db.flush()
        methodology_version = KnowledgeDocumentVersion(
            document_id=methodology.id,
            version_no=1,
            original_filename="method.md",
            content_type="text/markdown",
            source_type="methodology",
            sha256="c" * 64,
            storage_key="documents/method.md",
            status="published",
            effective_from=datetime.utcnow(),
        )
        case_version = KnowledgeDocumentVersion(
            document_id=case.id,
            version_no=1,
            original_filename="case.md",
            content_type="text/markdown",
            source_type="case_sop",
            sha256="d" * 64,
            storage_key="documents/case.md",
            status="published",
            effective_from=datetime.utcnow(),
        )
        db.add_all([methodology_version, case_version])
        await db.flush()
        keyword_match = KnowledgeChunk(
            version_id=methodology_version.id,
            chunk_no=1,
            content="流量 转化 方法",
            locator={"line_start": 2, "line_end": 4},
        )
        semantic_match = KnowledgeChunk(
            version_id=case_version.id,
            chunk_no=1,
            content="库存 周转 案例",
            locator={"line_start": 1, "line_end": 1},
        )
        db.add_all([keyword_match, semantic_match])
        await db.commit()

    service = KnowledgeRetrievalService(
        vector_index=RecalledVectorIndex([
            (semantic_match.id, 0.90),
            (keyword_match.id, 0.80),
        ]),
        embedding_service=FakeEmbedder(),
        session_factory=knowledge_db_factory,
    )
    service._ready = True

    results = await service.retrieve("流量 转化", {}, limit=2)

    assert [item.document_title for item in results] == ["转化方法", "经营案例"]
    assert [item.rank for item in results] == [1, 2]
    assert results[0].semantic_score == pytest.approx(0.80)
    assert results[0].keyword_score == pytest.approx(1.0)
    assert results[0].source_weight == pytest.approx(0.10)
    assert results[0].combined_score == pytest.approx(1.05)


@pytest.mark.asyncio
async def test_retrieval_boosts_heading_matches_for_chinese_queries(knowledge_db_factory):
    async with knowledge_db_factory() as db:
        first = KnowledgeDocument(title="Heading match", status="published")
        second = KnowledgeDocument(title="Content match", status="published")
        db.add_all([first, second])
        await db.flush()
        first_version = KnowledgeDocumentVersion(
            document_id=first.id,
            version_no=1,
            original_filename="heading.md",
            content_type="text/markdown",
            source_type="methodology",
            sha256="g" * 64,
            storage_key="documents/heading.md",
            status="published",
            effective_from=datetime.utcnow(),
        )
        second_version = KnowledgeDocumentVersion(
            document_id=second.id,
            version_no=1,
            original_filename="content.md",
            content_type="text/markdown",
            source_type="methodology",
            sha256="h" * 64,
            storage_key="documents/content.md",
            status="published",
            effective_from=datetime.utcnow(),
        )
        db.add_all([first_version, second_version])
        await db.flush()
        heading_match = KnowledgeChunk(
            version_id=first_version.id,
            chunk_no=1,
            content="general operations notes",
            locator={"heading_path": ["\u8f6c\u5316\u8bca\u65ad"]},
        )
        content_match = KnowledgeChunk(
            version_id=second_version.id,
            chunk_no=1,
            content="\u8f6c\u5316",
            locator={"heading_path": ["baseline"]},
        )
        db.add_all([heading_match, content_match])
        await db.commit()

    service = KnowledgeRetrievalService(
        vector_index=RecalledVectorIndex([
            (content_match.id, 0.8),
            (heading_match.id, 0.8),
        ]),
        embedding_service=FakeEmbedder(),
        session_factory=knowledge_db_factory,
    )
    service._ready = True

    results = await service.retrieve("\u8f6c\u5316\u8bca\u65ad", {}, limit=2)

    assert [item.chunk_id for item in results] == [heading_match.id, content_match.id]
    assert results[0].keyword_score == pytest.approx(0.35)
    assert results[1].keyword_score == pytest.approx(0.25)


@pytest.mark.asyncio
async def test_retrieval_strategies_control_scene_filtering_and_fallback(knowledge_db_factory):
    active_scene = {
        "industry": "retail",
        "sub_industry": "department-store",
        "business_mode": "self-operated",
        "operating_stage": "new",
    }
    async with knowledge_db_factory() as db:
        document = KnowledgeDocument(title="Retail replenishment", status="published")
        db.add(document)
        await db.flush()
        version = KnowledgeDocumentVersion(
            document_id=document.id,
            version_no=1,
            original_filename="retail.md",
            content_type="text/markdown",
            source_type="case_sop",
            sha256="e" * 64,
            storage_key="documents/retail.md",
            status="published",
            effective_from=datetime.utcnow(),
            industry_tags=["retail"],
        )
        db.add(version)
        await db.flush()
        chunk = KnowledgeChunk(
            version_id=version.id,
            chunk_no=1,
            content="inventory replenishment procedure",
            locator={},
        )
        db.add(chunk)
        await db.commit()

    vectors = RecordingVectorIndex([(chunk.id, 0.8)])
    service = KnowledgeRetrievalService(
        vector_index=vectors,
        embedding_service=FakeEmbedder(),
        session_factory=knowledge_db_factory,
    )
    service._ready = True

    assert await service.retrieve("inventory replenishment", active_scene, strategy_override="strict") == []
    industry_results = await service.retrieve(
        "inventory replenishment",
        active_scene,
        strategy_override="industry_only",
    )
    assert [item.chunk_id for item in industry_results] == [chunk.id]

    vectors.scenes.clear()
    progressive_results = await service.retrieve(
        "inventory replenishment",
        active_scene,
        strategy_override="progressive",
    )
    assert [item.chunk_id for item in progressive_results] == [chunk.id]
    assert vectors.scenes[-1] == {"industry": "retail"}

    unfiltered_results = await service.retrieve(
        "inventory replenishment",
        active_scene,
        strategy_override="unfiltered",
    )
    assert [item.chunk_id for item in unfiltered_results] == [chunk.id]
    assert vectors.scenes[-1] == {}

    boosted_results = await service.retrieve(
        "inventory replenishment",
        active_scene,
        strategy_override="scene_boost",
    )
    assert [item.chunk_id for item in boosted_results] == [chunk.id]
    assert boosted_results[0].combined_score == pytest.approx(
        unfiltered_results[0].combined_score + 0.04
    )
    assert vectors.scenes[-1] == {}

    async with knowledge_db_factory() as db:
        db.add(KnowledgeRetrievalConfiguration(id=1, strategy="industry_only"))
        await db.commit()
    global_results = await service.retrieve("inventory replenishment", active_scene)
    assert [item.chunk_id for item in global_results] == [chunk.id]


@pytest.mark.asyncio
async def test_report_citations_are_limited_to_authorized_candidates(knowledge_db_factory):
    async with knowledge_db_factory() as db:
        document = KnowledgeDocument(title="餐饮基准", status="published")
        db.add(document)
        await db.flush()
        version = KnowledgeDocumentVersion(
            document_id=document.id,
            version_no=1,
            original_filename="benchmark.txt",
            content_type="text/plain",
            source_type="benchmark_rule",
            sha256="b" * 64,
            storage_key="documents/version/benchmark.txt",
            status="published",
            effective_from=datetime.utcnow(),
            industry_tags=["餐饮"],
            sub_industry_tags=[],
            business_mode_tags=[],
            operating_stage_tags=[],
        )
        db.add(version)
        await db.flush()
        chunk = KnowledgeChunk(
            version_id=version.id,
            chunk_no=1,
            content="外卖转化异常应优先核对菜单曝光与下单漏斗。",
            locator={"line_start": 1, "line_end": 1},
        )
        db.add(chunk)
        session = await SessionRepository.for_system(db).create_session()
        report = await SessionRepository.for_system(db).save_report(session.id, "# 初始", {})
        await db.commit()

        candidate = EvidenceContext(
            chunk_id=chunk.id,
            version_id=version.id,
            document_id=document.id,
            document_title=document.title,
            source_type=version.source_type,
            version_no=version.version_no,
            quote=chunk.content,
            locator=chunk.locator,
            query="外卖转化",
            rank=1,
            retrieved_at=datetime.utcnow(),
        )
        markdown, selected = await validate_report_citations(
            db,
            "结论有来源。[证据 1] 伪造来源。[证据 9]",
            [candidate],
        )
        await persist_report_evidences(db, report_id=report.id, selected=selected)
        await db.commit()

        saved = list((await db.execute(
            select(ReportEvidence).where(ReportEvidence.report_id == report.id)
        )).scalars().all())

        version.status = "revoked"
        await db.flush()
        revoked_markdown, revoked_selected = await validate_report_citations(
            db,
            "撤回资料。[证据 1]",
            [candidate],
        )

    assert markdown == "结论有来源。[证据 1] 伪造来源。"
    assert len(saved) == 1
    assert revoked_markdown == "撤回资料。"
    assert revoked_selected == []
