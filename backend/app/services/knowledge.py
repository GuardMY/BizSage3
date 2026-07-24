"""Versioned industry knowledge ingestion, retrieval and evidence snapshots."""

import asyncio
import hashlib
import html
import io
import logging
import re
import uuid
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import PurePosixPath
from typing import Any, Iterable, Literal

import bleach
import markdown as markdown_renderer
import yaml
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from minio import Minio
from openai import AsyncOpenAI
from qdrant_client import AsyncQdrantClient, models as qmodels
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker
from sqlalchemy.orm import selectinload

from app.config import embedding_api_key, embedding_base_url, settings
from app.db import async_session_factory
from app.models import (
    KnowledgeAuditEvent,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeDocumentVersion,
    KnowledgeIngestionJob,
    KnowledgeRetrievalConfiguration,
    KnowledgeVectorSyncJob,
    ReportEvidence,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".docx", ".md", ".markdown", ".txt"}
SOURCE_TYPES = {"methodology", "benchmark_rule", "case_sop"}
RETRIEVAL_STRATEGIES = (
    "strict",
    "progressive",
    "industry_only",
    "unfiltered",
    "scene_boost",
)
RetrievalStrategy = Literal[
    "strict",
    "progressive",
    "industry_only",
    "unfiltered",
    "scene_boost",
]
_MIN_CHUNK_CHARS = 320
_TARGET_CHUNK_CHARS = 600
_MAX_CHUNK_CHARS = 800
_CHUNK_OVERLAP_CHARS = 80
_SCENE_FIELDS = ("industry", "sub_industry", "business_mode", "operating_stage")
_SCENE_MATCH_BOOST = 0.04


def utcnow() -> datetime:
    return datetime.utcnow()


def normalize_tag(value: str) -> str:
    return " ".join(value.strip().casefold().split())


def retrieval_terms(value: str) -> set[str]:
    """Tokenize mixed-language text, adding Chinese bigrams for precise matching."""
    terms: set[str] = set()
    for raw in re.findall(r"[\w\u4e00-\u9fff]+", value.casefold()):
        if len(raw) <= 1:
            continue
        terms.add(raw)
        if re.search(r"[\u4e00-\u9fff]", raw):
            terms.update(raw[index:index + 2] for index in range(len(raw) - 1))
    return terms


def normalize_retrieval_strategy(value: str | None) -> RetrievalStrategy | None:
    normalized = (value or "").strip().casefold()
    if normalized in RETRIEVAL_STRATEGIES:
        return normalized  # type: ignore[return-value]
    return None


def parse_tags(value: str | Iterable[str] | None) -> list[str]:
    """Normalize free-entered tags while preserving a deterministic order."""
    raw_items = re.split(r"[,，\n]", value) if isinstance(value, str) else (value or [])
    tags: list[str] = []
    for raw in raw_items:
        tag = normalize_tag(str(raw))
        if tag and tag not in tags:
            tags.append(tag)
    return tags


def extension_for(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def content_type_for(extension: str) -> str:
    return {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".txt": "text/plain",
    }[extension]


def safe_filename(filename: str) -> str:
    name = PurePosixPath(filename.replace("\\", "/")).name
    return name[:255] or "document"


@dataclass(frozen=True)
class ParsedBlock:
    text: str
    locator: dict[str, Any]


@dataclass(frozen=True)
class ParsedDocument:
    parser: str
    blocks: list[ParsedBlock]
    preview_html: str


@dataclass(frozen=True)
class EvidenceContext:
    """Authorized evidence given to a model or persisted with a report."""

    chunk_id: str
    version_id: str
    document_id: str
    document_title: str
    source_type: str
    version_no: int
    quote: str
    locator: dict[str, Any]
    query: str
    rank: int
    retrieved_at: datetime
    semantic_score: float = 0.0
    keyword_score: float = 0.0
    source_weight: float = 0.0
    combined_score: float = 0.0


def evidence_to_state(evidence: EvidenceContext) -> dict[str, Any]:
    """Convert evidence to JSON-only graph state for LangGraph checkpoints."""
    return {
        "chunk_id": evidence.chunk_id,
        "version_id": evidence.version_id,
        "document_id": evidence.document_id,
        "document_title": evidence.document_title,
        "source_type": evidence.source_type,
        "version_no": evidence.version_no,
        "quote": evidence.quote,
        "locator": evidence.locator,
        "query": evidence.query,
        "rank": evidence.rank,
        "retrieved_at": evidence.retrieved_at.isoformat(),
        "semantic_score": evidence.semantic_score,
        "keyword_score": evidence.keyword_score,
        "source_weight": evidence.source_weight,
        "combined_score": evidence.combined_score,
    }


def evidence_from_state(values: Iterable[EvidenceContext | dict[str, Any]]) -> list[EvidenceContext]:
    contexts: list[EvidenceContext] = []
    for value in values:
        if isinstance(value, EvidenceContext):
            contexts.append(value)
            continue
        try:
            retrieved_at = value.get("retrieved_at")
            contexts.append(EvidenceContext(
                chunk_id=str(value["chunk_id"]),
                version_id=str(value["version_id"]),
                document_id=str(value["document_id"]),
                document_title=str(value["document_title"]),
                source_type=str(value["source_type"]),
                version_no=int(value["version_no"]),
                quote=str(value["quote"]),
                locator=dict(value.get("locator") or {}),
                query=str(value["query"]),
                rank=int(value["rank"]),
                retrieved_at=datetime.fromisoformat(str(retrieved_at)),
                semantic_score=float(value.get("semantic_score", 0)),
                keyword_score=float(value.get("keyword_score", 0)),
                source_weight=float(value.get("source_weight", 0)),
                combined_score=float(value.get("combined_score", 0)),
            ))
        except (KeyError, TypeError, ValueError):
            logger.warning("Ignoring malformed evidence in graph state")
    return contexts


class KnowledgeStorage:
    """MinIO wrapper. Object access is always proxied through authorized APIs."""

    def __init__(self) -> None:
        self._client = Minio(
            settings.minio_endpoint,
            access_key=settings.minio_access_key,
            secret_key=settings.minio_secret_key,
            secure=settings.minio_secure,
        )

    async def startup(self) -> None:
        await self._retry(self._ensure_bucket)

    async def _retry(self, operation, attempts: int = 15) -> None:
        last_error: Exception | None = None
        for attempt in range(attempts):
            try:
                await asyncio.to_thread(operation)
                return
            except Exception as exc:  # MinIO may start after FastAPI.
                last_error = exc
                await asyncio.sleep(min(0.5 * (attempt + 1), 3))
        raise RuntimeError("MinIO is unavailable") from last_error

    def _ensure_bucket(self) -> None:
        if not self._client.bucket_exists(settings.minio_bucket):
            self._client.make_bucket(settings.minio_bucket)

    async def put(self, storage_key: str, content: bytes, content_type: str) -> None:
        def operation() -> None:
            self._client.put_object(
                settings.minio_bucket,
                storage_key,
                io.BytesIO(content),
                len(content),
                content_type=content_type,
            )

        await asyncio.to_thread(operation)

    async def get(self, storage_key: str) -> bytes:
        def operation() -> bytes:
            response = self._client.get_object(settings.minio_bucket, storage_key)
            try:
                return response.read()
            finally:
                response.close()
                response.release_conn()

        return await asyncio.to_thread(operation)

    async def delete(self, storage_key: str) -> None:
        await asyncio.to_thread(
            self._client.remove_object,
            settings.minio_bucket,
            storage_key,
        )

    async def clear(self) -> None:
        def operation() -> None:
            for item in self._client.list_objects(settings.minio_bucket, recursive=True):
                self._client.remove_object(settings.minio_bucket, item.object_name)

        await asyncio.to_thread(operation)


class EmbeddingService:
    def __init__(self) -> None:
        self._client = AsyncOpenAI(
            api_key=embedding_api_key(),
            base_url=embedding_base_url(),
        )

    async def embed(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        if not embedding_api_key():
            raise RuntimeError("未配置 OPENAI_API_KEY 或 EMBEDDING_API_KEY")
        response = await self._client.embeddings.create(
            model=settings.embedding_model,
            input=texts,
            dimensions=settings.embedding_dimensions,
        )
        return [item.embedding for item in response.data]


class KnowledgeVectorIndex:
    def __init__(self) -> None:
        self._client = AsyncQdrantClient(url=settings.qdrant_url)

    async def startup(self) -> None:
        async def ensure_collection() -> None:
            collections = await self._client.get_collections()
            if not any(item.name == settings.qdrant_collection for item in collections.collections):
                await self._client.create_collection(
                    collection_name=settings.qdrant_collection,
                    vectors_config=qmodels.VectorParams(
                        size=settings.embedding_dimensions,
                        distance=qmodels.Distance.COSINE,
                    ),
                )

        last_error: Exception | None = None
        for attempt in range(15):
            try:
                await ensure_collection()
                return
            except Exception as exc:
                last_error = exc
                await asyncio.sleep(min(0.5 * (attempt + 1), 3))
        raise RuntimeError("Qdrant is unavailable") from last_error

    async def close(self) -> None:
        await self._client.close()

    async def reset_collection(self) -> None:
        collections = await self._client.get_collections()
        if any(item.name == settings.qdrant_collection for item in collections.collections):
            await self._client.delete_collection(settings.qdrant_collection)
        await self.startup()

    async def upsert_chunks(
        self,
        chunks: list[KnowledgeChunk],
        vectors: list[list[float]],
        version: KnowledgeDocumentVersion,
    ) -> None:
        points: list[qmodels.PointStruct] = []
        for chunk, vector in zip(chunks, vectors, strict=True):
            point_id = chunk.embedding_ref or str(
                uuid.uuid5(uuid.NAMESPACE_URL, f"bizsage:{chunk.id}")
            )
            chunk.embedding_ref = point_id
            points.append(
                qmodels.PointStruct(
                    id=point_id,
                    vector=vector,
                    payload={
                        "chunk_id": chunk.id,
                        "version_id": version.id,
                        "status": version.status,
                        "industry_tags": version.industry_tags or [],
                        "sub_industry_tags": version.sub_industry_tags or [],
                        "business_mode_tags": version.business_mode_tags or [],
                        "operating_stage_tags": version.operating_stage_tags or [],
                        "source_type": version.source_type,
                        "chunk_kind": str((chunk.locator or {}).get("kind", "paragraph")),
                        "heading_path": [
                            str(item)
                            for item in (chunk.locator or {}).get("heading_path", [])
                        ],
                    },
                )
            )
        if points:
            await self._client.upsert(
                collection_name=settings.qdrant_collection,
                points=points,
                wait=True,
            )

    async def activate_version(self, version_id: str) -> None:
        await self._client.set_payload(
            collection_name=settings.qdrant_collection,
            payload={"status": "published"},
            points=qmodels.FilterSelector(
                filter=qmodels.Filter(
                    must=[qmodels.FieldCondition(
                        key="version_id", match=qmodels.MatchValue(value=version_id)
                    )]
                )
            ),
            wait=True,
        )

    async def delete_version(self, version_id: str) -> None:
        await self._client.delete(
            collection_name=settings.qdrant_collection,
            points_selector=qmodels.FilterSelector(
                filter=qmodels.Filter(
                    must=[qmodels.FieldCondition(
                        key="version_id", match=qmodels.MatchValue(value=version_id)
                    )]
                )
            ),
            wait=True,
        )

    async def search(
        self,
        vector: list[float],
        scene: dict[str, str],
        limit: int,
    ) -> list[tuple[str, float]]:
        filters = [
            qmodels.FieldCondition(
                key="status", match=qmodels.MatchValue(value="published")
            )
        ]
        metadata_fields = {
            "industry": "industry_tags",
            "sub_industry": "sub_industry_tags",
            "business_mode": "business_mode_tags",
            "operating_stage": "operating_stage_tags",
        }
        for scene_key, payload_key in metadata_fields.items():
            value = normalize_tag(scene.get(scene_key, ""))
            if value:
                filters.append(
                    qmodels.FieldCondition(
                        key=payload_key,
                        match=qmodels.MatchValue(value=value),
                    )
                )
        result = await self._client.query_points(
            collection_name=settings.qdrant_collection,
            query=vector,
            query_filter=qmodels.Filter(must=filters),
            limit=limit,
            with_payload=["chunk_id"],
            with_vectors=False,
        )
        return [
            (str(point.payload["chunk_id"]), float(point.score))
            for point in result.points
            if point.payload and point.payload.get("chunk_id")
        ]


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    extension = extension_for(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 DOCX、Markdown 和 TXT 文件")
    if extension == ".docx":
        return _parse_docx(content)
    return _parse_text(content, extension)


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文本文件必须使用 UTF-8 或 GB18030 编码")


_MARKDOWN_TABLE_SEPARATOR = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)
_MARKDOWN_LIST_ITEM = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)")
_SENTENCE_BOUNDARIES = ("\n", "。", "！", "？", "；", ".", "!", "?", ";")


def _strip_markdown_frontmatter(raw: str) -> str:
    """Blank a leading YAML frontmatter block while preserving source line numbers."""
    lines = raw.split("\n")
    if not lines or lines[0].strip() != "---":
        return raw
    try:
        end = next(
            index for index, line in enumerate(lines[1:], start=1)
            if line.strip() == "---"
        )
    except StopIteration:
        return raw
    try:
        frontmatter = yaml.safe_load("\n".join(lines[1:end]))
    except yaml.YAMLError:
        return raw
    if not isinstance(frontmatter, dict):
        return raw
    lines[:end + 1] = [""] * (end + 1)
    return "\n".join(lines)


def _is_markdown_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return "|" in lines[index] and bool(_MARKDOWN_TABLE_SEPARATOR.match(lines[index + 1]))


def _parse_text(content: bytes, extension: str) -> ParsedDocument:
    raw = _decode_text(content).replace("\r\n", "\n").replace("\r", "\n")
    body = _strip_markdown_frontmatter(raw) if extension != ".txt" else raw
    lines = body.split("\n")
    heading_path: list[str] = []
    blocks: list[ParsedBlock] = []
    table_no = 0
    index = 0
    while index < len(lines):
        line = lines[index]
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line) if extension != ".txt" else None
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_path[:] = heading_path[: level - 1]
            heading_path.append(title)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if extension != ".txt" and _is_markdown_table_start(lines, index):
            start = index
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                table_lines.append(lines[index].strip())
                index += 1
            if len(table_lines) >= 2:
                table_no += 1
                blocks.append(ParsedBlock(
                    text="\n".join(table_lines),
                    locator={
                        "heading_path": list(heading_path),
                        "kind": "table",
                        "table_format": "markdown",
                        "table_no": table_no,
                        "row_start": 1,
                        "row_end": max(1, len(table_lines) - 2),
                        "line_start": start + 1,
                        "line_end": index,
                    },
                ))
                continue
            index = start

        start = index
        paragraph: list[str] = []
        while index < len(lines):
            candidate = lines[index]
            if not candidate.strip():
                break
            if extension != ".txt" and re.match(r"^(#{1,6})\s+(.+?)\s*$", candidate):
                break
            if extension != ".txt" and _is_markdown_table_start(lines, index):
                break
            paragraph.append(candidate.strip())
            index += 1
        text = "\n".join(paragraph).strip()
        if text:
            blocks.append(ParsedBlock(
                text=text,
                locator={
                    "heading_path": list(heading_path),
                    "kind": "list" if any(_MARKDOWN_LIST_ITEM.match(item) for item in paragraph) else "paragraph",
                    "line_start": start + 1,
                    "line_end": index,
                },
            ))
        elif index == start:
            index += 1
    if not blocks and body.strip():
        blocks = [ParsedBlock(
            body.strip(),
            {"kind": "paragraph", "line_start": 1, "line_end": raw.count("\n") + 1},
        )]

    if extension == ".txt":
        preview_html = f"<pre>{html.escape(raw)}</pre>"
    else:
        rendered = markdown_renderer.markdown(body, extensions=["extra", "sane_lists"])
        preview_html = bleach.clean(
            rendered,
            tags=["p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "strong", "em", "code", "pre", "blockquote", "table", "thead", "tbody", "tr", "th", "td", "a", "br"],
            attributes={"a": ["href", "title"]},
            protocols=["http", "https", "mailto"],
            strip=True,
        )
    return ParsedDocument(parser="markdown" if extension != ".txt" else "text", blocks=blocks, preview_html=preview_html)


def _parse_docx(content: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("无法读取 DOCX 文件") from exc

    heading_path: list[str] = []
    blocks: list[ParsedBlock] = []
    preview: list[str] = []
    paragraph_no = 0
    table_no = 0
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_no += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (getattr(paragraph.style, "name", "") or "").lower()
            match = re.search(r"(?:heading|标题)\s*(\d+)", style_name)
            if match:
                level = int(match.group(1))
                heading_path[:] = heading_path[: level - 1]
                heading_path.append(text)
                preview.append(f"<h{min(level, 6)}>{html.escape(text)}</h{min(level, 6)}>")
                continue
            blocks.append(ParsedBlock(
                text=text,
                locator={
                    "heading_path": list(heading_path),
                    "kind": "paragraph",
                    "paragraph_start": paragraph_no,
                    "paragraph_end": paragraph_no,
                },
            ))
            preview.append(f"<p>{html.escape(text)}</p>")
            continue
        if not child.tag.endswith("}tbl"):
            continue
        table_no += 1
        table = Table(child, document)
        rows = [" | ".join(cell.text.strip() for cell in row.cells).strip() for row in table.rows]
        table_text = "\n".join(row for row in rows if row)
        if table_text:
            blocks.append(ParsedBlock(
                text=table_text,
                locator={
                    "heading_path": list(heading_path),
                    "kind": "table",
                    "table_format": "docx",
                    "table_no": table_no,
                    "row_start": 1,
                    "row_end": max(1, len(rows) - 1),
                },
            ))
            preview.append(f"<pre>{html.escape(table_text)}</pre>")

    if not blocks:
        raise ValueError("DOCX 中没有可解析的正文内容")
    return ParsedDocument(parser="docx", blocks=blocks, preview_html="\n".join(preview))


def _section_key(locator: dict[str, Any]) -> tuple[str, ...]:
    heading_path = locator.get("heading_path")
    if not isinstance(heading_path, list):
        return ()
    return tuple(str(item) for item in heading_path if str(item))


def _split_prose(text: str) -> list[str]:
    """Split long prose at a natural boundary and retain a small local overlap."""
    remaining = text.strip()
    parts: list[str] = []
    while len(remaining) > _MAX_CHUNK_CHARS:
        candidate = remaining[:_MAX_CHUNK_CHARS]
        boundaries = [
            candidate.rfind(boundary, _MIN_CHUNK_CHARS)
            for boundary in _SENTENCE_BOUNDARIES
        ]
        end = max(boundaries) + 1
        if end <= _MIN_CHUNK_CHARS:
            end = _MAX_CHUNK_CHARS
        parts.append(remaining[:end].strip())
        overlap = remaining[max(0, end - _CHUNK_OVERLAP_CHARS):end].strip()
        remaining = f"{overlap}{remaining[end:]}".strip()
    if remaining:
        parts.append(remaining)
    return parts


def _table_chunks(block: ParsedBlock) -> list[tuple[str, dict[str, Any]]]:
    """Keep Markdown table rows independently citable and repeat headers when split."""
    if len(block.text) <= _MAX_CHUNK_CHARS:
        return [(block.text, dict(block.locator))]
    lines = block.text.split("\n")
    if len(lines) < 2:
        return [(part, dict(block.locator)) for part in _split_prose(block.text)]
    if block.locator.get("table_format") == "markdown":
        if len(lines) < 3:
            return [(part, dict(block.locator)) for part in _split_prose(block.text)]
        prefix = lines[:2]
        rows = lines[2:]
    else:
        prefix = lines[:1]
        rows = lines[1:]
    chunks: list[tuple[str, dict[str, Any]]] = []
    current_rows: list[str] = []
    row_start = 1

    def flush_rows() -> None:
        nonlocal current_rows, row_start
        if not current_rows:
            return
        locator = dict(block.locator)
        locator["row_start"] = row_start
        locator["row_end"] = row_start + len(current_rows) - 1
        chunks.append(("\n".join([*prefix, *current_rows]), locator))
        row_start += len(current_rows)
        current_rows = []

    for row in rows:
        projected = "\n".join([*prefix, *current_rows, row])
        if current_rows and len(projected) > _MAX_CHUNK_CHARS:
            flush_rows()
        if len("\n".join([*prefix, row])) > _MAX_CHUNK_CHARS:
            for part in _split_prose(row):
                locator = dict(block.locator)
                locator["row_start"] = row_start
                locator["row_end"] = row_start
                chunks.append(("\n".join([*prefix, part]), locator))
            row_start += 1
            continue
        current_rows.append(row)
    flush_rows()
    return chunks


def _merged_locator(first: dict[str, Any], last: dict[str, Any]) -> dict[str, Any]:
    locator = dict(first)
    if first.get("kind") != last.get("kind"):
        locator["kind"] = "section"
    for key in ("line_end", "paragraph_end"):
        if key in last:
            locator[key] = last[key]
    return locator


def chunk_blocks(blocks: list[ParsedBlock]) -> list[tuple[str, dict[str, Any]]]:
    """Create heading-scoped, citation-friendly chunks without cross-section overlap."""
    chunks: list[tuple[str, dict[str, Any]]] = []
    buffer: list[str] = []
    first_locator: dict[str, Any] | None = None
    last_locator: dict[str, Any] | None = None
    active_section: tuple[str, ...] | None = None

    def flush() -> None:
        nonlocal buffer, first_locator, last_locator, active_section
        if buffer and first_locator is not None and last_locator is not None:
            chunks.append(("\n\n".join(buffer).strip(), _merged_locator(first_locator, last_locator)))
        buffer = []
        first_locator = None
        last_locator = None
        active_section = None

    for block in blocks:
        text = block.text.strip()
        if not text:
            continue
        section = _section_key(block.locator)
        if buffer and section != active_section:
            flush()
        if block.locator.get("kind") == "table":
            flush()
            chunks.extend(_table_chunks(block))
            continue
        for part in _split_prose(text):
            current_size = len("\n\n".join(buffer))
            projected = len("\n\n".join([*buffer, part]))
            if buffer and (
                projected > _MAX_CHUNK_CHARS
                or (projected > _TARGET_CHUNK_CHARS and current_size >= _MIN_CHUNK_CHARS)
            ):
                flush()
            if not buffer:
                first_locator = dict(block.locator)
                active_section = section
            buffer.append(part)
            last_locator = dict(block.locator)
    flush()
    return chunks


def embedding_text_for(
    document_title: str,
    content: str,
    locator: dict[str, Any],
) -> str:
    """Give embeddings title and section context while keeping citations verbatim."""
    heading_path = locator.get("heading_path")
    headings = " / ".join(
        str(item).strip()
        for item in heading_path
        if str(item).strip()
    ) if isinstance(heading_path, list) else ""
    context = [item for item in (document_title.strip(), headings) if item]
    return "\n".join([*context, content])


def _metadata_matches(version: KnowledgeDocumentVersion, scene: dict[str, str]) -> bool:
    for key, tags in (
        ("industry", version.industry_tags),
        ("sub_industry", version.sub_industry_tags),
        ("business_mode", version.business_mode_tags),
        ("operating_stage", version.operating_stage_tags),
    ):
        expected = normalize_tag(scene.get(key, ""))
        if expected and expected not in (tags or []):
            return False
    return True


def _scene_subset(scene: dict[str, str], fields: tuple[str, ...]) -> dict[str, str]:
    return {
        field: str(scene.get(field) or "")
        for field in fields
        if scene.get(field)
    }


def _filter_scenes_for_strategy(
    scene: dict[str, str],
    strategy: RetrievalStrategy,
) -> list[dict[str, str]]:
    full_scene = _scene_subset(scene, _SCENE_FIELDS)
    if strategy == "strict":
        return [full_scene]
    if strategy == "industry_only":
        return [_scene_subset(scene, ("industry",))]
    if strategy in {"unfiltered", "scene_boost"}:
        return [{}]

    candidates = [
        full_scene,
        _scene_subset(scene, ("industry", "sub_industry", "business_mode")),
        _scene_subset(scene, ("industry", "sub_industry")),
        _scene_subset(scene, ("industry",)),
        {},
    ]
    unique: list[dict[str, str]] = []
    for candidate in candidates:
        if candidate not in unique:
            unique.append(candidate)
    return unique


def _scene_match_count(version: KnowledgeDocumentVersion, scene: dict[str, str]) -> int:
    matches = 0
    for key, tags in (
        ("industry", version.industry_tags),
        ("sub_industry", version.sub_industry_tags),
        ("business_mode", version.business_mode_tags),
        ("operating_stage", version.operating_stage_tags),
    ):
        expected = normalize_tag(scene.get(key, ""))
        if expected and expected in (tags or []):
            matches += 1
    return matches


class KnowledgeRetrievalService:
    """Retrieval always validates relational version state after Qdrant recall."""

    def __init__(
        self,
        *,
        vector_index: KnowledgeVectorIndex | None = None,
        embedding_service: EmbeddingService | None = None,
        session_factory: async_sessionmaker[AsyncSession] = async_session_factory,
    ) -> None:
        self._vector_index = vector_index or KnowledgeVectorIndex()
        self._embedding_service = embedding_service or EmbeddingService()
        self._session_factory = session_factory
        self._ready = False

    async def startup(self) -> None:
        await self._vector_index.startup()
        self._ready = True

    async def shutdown(self) -> None:
        await self._vector_index.close()
        self._ready = False

    @staticmethod
    def _environment_default_strategy() -> RetrievalStrategy:
        return normalize_retrieval_strategy(
            settings.knowledge_retrieval_default_strategy
        ) or "strict"

    async def get_global_strategy(self) -> RetrievalStrategy:
        try:
            async with self._session_factory() as db:
                configuration = await db.get(KnowledgeRetrievalConfiguration, 1)
        except Exception:
            logger.exception("Knowledge retrieval policy lookup failed; using environment default")
            return self._environment_default_strategy()
        return normalize_retrieval_strategy(
            configuration.strategy if configuration else None
        ) or self._environment_default_strategy()

    async def set_global_strategy(self, strategy: str) -> RetrievalStrategy:
        normalized = normalize_retrieval_strategy(strategy)
        if normalized is None:
            raise ValueError("invalid_retrieval_strategy")
        async with self._session_factory() as db:
            configuration = await db.get(KnowledgeRetrievalConfiguration, 1)
            if configuration is None:
                db.add(KnowledgeRetrievalConfiguration(id=1, strategy=normalized))
            else:
                configuration.strategy = normalized
            await db.commit()
        return normalized

    async def retrieve(
        self,
        query: str,
        scene: dict[str, str],
        *,
        limit: int = 5,
        raise_on_error: bool = False,
        strategy_override: str | None = None,
    ) -> list[EvidenceContext]:
        if not self._ready or not query.strip():
            return []
        if strategy_override is not None:
            strategy = normalize_retrieval_strategy(strategy_override)
            if strategy is None:
                raise ValueError("invalid_retrieval_strategy")
        else:
            strategy = await self.get_global_strategy()
        try:
            vector = (await self._embedding_service.embed([query]))[0]
        except Exception:
            logger.exception("Knowledge retrieval failed")
            if raise_on_error:
                raise
            return []

        for filter_scene in _filter_scenes_for_strategy(scene, strategy):
            try:
                recalled = await self._vector_index.search(
                    vector,
                    filter_scene,
                    max(limit * 4, 12),
                )
                selected = await self._rank_recalled(
                    recalled,
                    query=query,
                    filter_scene=filter_scene,
                    ranking_scene=scene if strategy == "scene_boost" else {},
                    limit=limit,
                )
            except Exception:
                logger.exception("Knowledge retrieval failed")
                if raise_on_error:
                    raise
                return []
            if selected:
                return selected
        return []

    async def _rank_recalled(
        self,
        recalled: list[tuple[str, float]],
        *,
        query: str,
        filter_scene: dict[str, str],
        ranking_scene: dict[str, str],
        limit: int,
    ) -> list[EvidenceContext]:
        if not recalled:
            return []
        chunk_ids = [chunk_id for chunk_id, _ in recalled]
        now = utcnow()
        async with self._session_factory() as db:
            stmt = (
                select(KnowledgeChunk)
                .join(KnowledgeDocumentVersion)
                .join(KnowledgeDocument)
                .where(
                    KnowledgeChunk.id.in_(chunk_ids),
                    KnowledgeDocumentVersion.status == "published",
                    KnowledgeDocumentVersion.effective_from <= now,
                    (KnowledgeDocumentVersion.effective_to.is_(None) | (KnowledgeDocumentVersion.effective_to > now)),
                )
                .options(selectinload(KnowledgeChunk.version).selectinload(KnowledgeDocumentVersion.document))
            )
            db_chunks = {chunk.id: chunk for chunk in (await db.execute(stmt)).scalars().all()}

        query_terms = retrieval_terms(query)
        ranked: list[tuple[float, EvidenceContext]] = []
        source_weight = {"methodology": 0.10, "benchmark_rule": 0.08, "case_sop": 0.05}
        for rank, (chunk_id, semantic_score) in enumerate(recalled, start=1):
            chunk = db_chunks.get(chunk_id)
            if not chunk or not _metadata_matches(chunk.version, filter_scene):
                continue
            text_terms = retrieval_terms(chunk.content)
            heading_terms = retrieval_terms(" ".join(
                str(item) for item in (chunk.locator or {}).get("heading_path", [])
            ))
            content_match = len(query_terms & text_terms) / max(len(query_terms), 1)
            heading_match = len(query_terms & heading_terms) / max(len(query_terms), 1)
            keyword_score = min(1.0, content_match + heading_match * 0.35)
            type_weight = source_weight.get(chunk.version.source_type, 0)
            scene_boost = _scene_match_count(chunk.version, ranking_scene) * _SCENE_MATCH_BOOST
            score = semantic_score + keyword_score * 0.15 + type_weight + scene_boost
            ranked.append((score, EvidenceContext(
                chunk_id=chunk.id,
                version_id=chunk.version.id,
                document_id=chunk.version.document.id,
                document_title=chunk.version.document.title,
                source_type=chunk.version.source_type,
                version_no=chunk.version.version_no,
                quote=chunk.content,
                locator=chunk.locator or {},
                query=query,
                rank=rank,
                retrieved_at=now,
                semantic_score=semantic_score,
                keyword_score=keyword_score,
                source_weight=type_weight,
                combined_score=score,
            )))
        ranked.sort(key=lambda item: item[0], reverse=True)
        return [
            replace(item[1], rank=rank)
            for rank, item in enumerate(ranked[:limit], start=1)
        ]


_CITATION_PATTERN = re.compile(r"\[证据\s*(\d+)\]")


async def validate_report_citations(
    db: AsyncSession,
    markdown: str,
    candidates: list[EvidenceContext],
) -> tuple[str, list[tuple[int, EvidenceContext]]]:
    """Keep only citations backed by this retrieval and a non-revoked source.

    The model sees no storage URLs or raw identifiers. This server-side check is
    the boundary that prevents it from creating arbitrary source references.
    """
    candidate_by_no = {index: evidence for index, evidence in enumerate(candidates, start=1)}
    requested = {int(match.group(1)) for match in _CITATION_PATTERN.finditer(markdown)}
    if not requested:
        return markdown, []
    selected = {number: candidate_by_no[number] for number in requested if number in candidate_by_no}
    if selected:
        valid_rows = (await db.execute(
            select(KnowledgeChunk.id, KnowledgeChunk.version_id)
            .join(KnowledgeDocumentVersion)
            .where(
                KnowledgeChunk.id.in_([evidence.chunk_id for evidence in selected.values()]),
                KnowledgeDocumentVersion.status != "revoked",
            )
        )).all()
        valid_pairs = {(row[0], row[1]) for row in valid_rows}
        selected = {
            number: evidence
            for number, evidence in selected.items()
            if (evidence.chunk_id, evidence.version_id) in valid_pairs
        }

    valid_numbers = set(selected)
    sanitized = _CITATION_PATTERN.sub(
        lambda match: match.group(0) if int(match.group(1)) in valid_numbers else "",
        markdown,
    )
    return sanitized, sorted(selected.items())


async def persist_report_evidences(
    db: AsyncSession,
    *,
    report_id: str,
    selected: list[tuple[int, EvidenceContext]],
) -> None:
    for evidence_no, evidence in selected:
        db.add(ReportEvidence(
            report_id=report_id,
            version_id=evidence.version_id,
            chunk_id=evidence.chunk_id,
            evidence_no=evidence_no,
            quote=evidence.quote,
            locator=evidence.locator,
            query=evidence.query,
            rank=evidence.rank,
            retrieved_at=evidence.retrieved_at,
        ))
        db.add(KnowledgeAuditEvent(
            event_type="report_evidence_persisted",
            document_id=evidence.document_id,
            version_id=evidence.version_id,
            actor_role="system",
            detail={"report_id": report_id, "chunk_id": evidence.chunk_id, "evidence_no": evidence_no},
        ))
    await db.flush()


class KnowledgeIngestionService:
    """Execute durable ingestion jobs without holding transactions over I/O."""

    def __init__(
        self,
        *,
        storage: KnowledgeStorage | None = None,
        vector_index: KnowledgeVectorIndex | None = None,
        embedding_service: EmbeddingService | None = None,
        session_factory: async_sessionmaker[AsyncSession] = async_session_factory,
    ) -> None:
        self._storage = storage or knowledge_storage
        self._vector_index = vector_index or knowledge_vector_index
        self._embedding_service = embedding_service or globals().get("embedding_service") or EmbeddingService()
        self._session_factory = session_factory
        self._max_attempts = settings.background_job_max_attempts

    async def run(self, job_id: str, *, worker_id: str) -> bool:
        claimed = await self._claim(job_id, worker_id)
        if claimed is None:
            return False
        version, job, attempt_count = claimed
        try:
            content = await self._storage.get(version.storage_key)
            parsed = parse_document(version.original_filename, content)
            chunk_data = chunk_blocks(parsed.blocks)
            if not chunk_data:
                raise ValueError("未提取到可索引的正文内容")
            async with self._session_factory() as db:
                document = await db.get(KnowledgeDocument, version.document_id)
                if document is None:
                    raise ValueError("知识资料不存在")
                document_title = document.title
            vectors = await self._embedding_service.embed([
                embedding_text_for(document_title, text, locator)
                for text, locator in chunk_data
            ])
            if len(vectors) != len(chunk_data):
                raise RuntimeError("嵌入服务返回的向量数量不正确")

            chunks = [
                KnowledgeChunk(
                    id=uuid.uuid4().hex,
                    version_id=version.id,
                    chunk_no=index,
                    content=text,
                    locator=locator,
                )
                for index, (text, locator) in enumerate(chunk_data, start=1)
            ]
            for chunk in chunks:
                chunk.embedding_ref = str(
                    uuid.uuid5(uuid.NAMESPACE_URL, f"bizsage:{chunk.id}")
                )

            async with self._session_factory() as db:
                current_version = await db.get(KnowledgeDocumentVersion, version.id)
                current_job = await db.get(KnowledgeIngestionJob, job.id)
                if (
                    current_version is None
                    or current_job is None
                    or current_job.state != "running"
                    or current_version.status == "revoked"
                    or current_job.worker_id != worker_id
                ):
                    return False
                await db.execute(
                    delete(KnowledgeChunk).where(KnowledgeChunk.version_id == version.id)
                )
                db.add_all(chunks)
                current_version.status = "indexing"
                await db.commit()

            version.status = "indexing"
            await self._vector_index.upsert_chunks(chunks, vectors, version)

            async with self._session_factory() as db:
                current_version = await db.get(KnowledgeDocumentVersion, version.id)
                current_job = await db.get(KnowledgeIngestionJob, job.id)
                if (
                    current_version is None
                    or current_job is None
                    or current_job.state != "running"
                    or current_job.worker_id != worker_id
                ):
                    return False
                if current_version.status == "revoked":
                    current_job.state = "cancelled"
                    current_job.finished_at = utcnow()
                    await db.commit()
                    await self._vector_index.delete_version(version.id)
                    return False
                current_version.status = "pending_review"
                current_job.state = "completed"
                current_job.parser = parsed.parser
                current_job.indexed_at = utcnow()
                current_job.finished_at = utcnow()
                current_job.worker_id = worker_id
                current_job.error = None
                db.add(KnowledgeAuditEvent(
                    event_type="ingestion_completed",
                    document_id=current_version.document_id,
                    version_id=current_version.id,
                    actor_role="system",
                    detail={"parser": parsed.parser, "chunk_count": len(chunks)},
                ))
                document = await db.get(KnowledgeDocument, current_version.document_id)
                if document is not None and document.managed_source_key:
                    from app.services.industry_catalog import mark_catalog_item_finished
                    from app.services.knowledge_lifecycle import publish_version

                    await publish_version(db, current_version, actor_role="system")
                    await mark_catalog_item_finished(
                        db,
                        current_job.id,
                        state="published",
                    )
                await db.commit()
            logger.info("Knowledge ingestion completed: job_id=%s version_id=%s", job.id, version.id)
            return True
        except asyncio.CancelledError:
            try:
                await self._vector_index.delete_version(version.id)
            except Exception:
                logger.exception("Failed to clean cancelled Qdrant version %s", version.id)
            await self._record_failure(job.id, version.id, attempt_count, worker_id, "知识入库任务被中断")
            raise
        except Exception as exc:
            logger.exception("Knowledge ingestion failed: job_id=%s version_id=%s", job.id, version.id)
            try:
                await self._vector_index.delete_version(version.id)
            except Exception:
                logger.exception("Failed to clean partial Qdrant version %s", version.id)
            await self._record_failure(
                job.id,
                version.id,
                attempt_count,
                worker_id,
                str(exc)[:2000],
            )
            raise

    async def _claim(
        self,
        job_id: str,
        worker_id: str,
    ) -> tuple[KnowledgeDocumentVersion, KnowledgeIngestionJob, int] | None:
        now = utcnow()
        async with self._session_factory() as db:
            result = await db.execute(
                update(KnowledgeIngestionJob)
                .where(
                    KnowledgeIngestionJob.id == job_id,
                    KnowledgeIngestionJob.state == "queued",
                )
                .values(
                    state="running",
                    worker_id=worker_id,
                    started_at=now,
                    finished_at=None,
                    error=None,
                )
                .returning(KnowledgeIngestionJob.version_id, KnowledgeIngestionJob.retry_count)
            )
            row = result.one_or_none()
            if row is None:
                await db.rollback()
                return None
            version = await db.get(KnowledgeDocumentVersion, row.version_id)
            job = await db.get(KnowledgeIngestionJob, job_id)
            if version is None or job is None or version.status == "revoked":
                if job is not None:
                    job.state = "cancelled"
                    job.finished_at = now
                await db.commit()
                return None
            version.status = "parsing"
            from app.services.industry_catalog import mark_catalog_item_running

            await mark_catalog_item_running(db, job_id)
            await db.commit()
            return version, job, int(row.retry_count) + 1

    async def _record_failure(
        self,
        job_id: str,
        version_id: str,
        attempt_count: int,
        worker_id: str,
        error: str,
    ) -> None:
        retrying = attempt_count < self._max_attempts
        async with self._session_factory() as db:
            version = await db.get(KnowledgeDocumentVersion, version_id)
            job = await db.get(KnowledgeIngestionJob, job_id)
            if job is None or job.state != "running" or job.worker_id != worker_id:
                return
            if version is not None and version.status != "revoked":
                version.status = "draft"
            job.state = "queued" if retrying else "failed"
            job.error = error
            job.retry_count += 1
            job.worker_id = None
            job.finished_at = None if retrying else utcnow()
            if version is not None:
                db.add(KnowledgeAuditEvent(
                    event_type="ingestion_failed",
                    document_id=version.document_id,
                    version_id=version.id,
                    actor_role="system",
                    detail={"error": error[:500], "will_retry": retrying},
                ))
            from app.services.industry_catalog import mark_catalog_item_finished

            await mark_catalog_item_finished(
                db,
                job.id,
                state="failed",
                error=error,
                retrying=retrying,
            )
            await db.commit()


class KnowledgeVectorSyncService:
    """Execute durable, idempotent Qdrant activation/deletion jobs."""

    def __init__(
        self,
        *,
        vector_index: KnowledgeVectorIndex | None = None,
        session_factory: async_sessionmaker[AsyncSession] = async_session_factory,
    ) -> None:
        self._vector_index = vector_index or knowledge_vector_index
        self._session_factory = session_factory
        self._max_attempts = settings.background_job_max_attempts

    async def run(self, job_id: str, *, worker_id: str) -> bool:
        claimed = await self._claim(job_id, worker_id)
        if claimed is None:
            return False
        version_id, operation, attempt_count = claimed
        try:
            if operation == "activate":
                await self._vector_index.activate_version(version_id)
            elif operation == "delete":
                await self._vector_index.delete_version(version_id)
            else:
                raise ValueError(f"Unsupported vector sync operation: {operation}")
            async with self._session_factory() as db:
                job = await db.get(KnowledgeVectorSyncJob, job_id)
                if job is None or job.state != "running" or job.worker_id != worker_id:
                    return False
                job.state = "completed"
                job.error = None
                job.finished_at = utcnow()
                await db.commit()
            return True
        except Exception as exc:
            logger.exception("Knowledge vector sync failed: %s", job_id)
            await self._record_failure(job_id, version_id, attempt_count, worker_id, str(exc)[:2000])
            raise

    async def _claim(self, job_id: str, worker_id: str) -> tuple[str, str, int] | None:
        async with self._session_factory() as db:
            result = await db.execute(
                update(KnowledgeVectorSyncJob)
                .where(
                    KnowledgeVectorSyncJob.id == job_id,
                    KnowledgeVectorSyncJob.state == "queued",
                )
                .values(
                    state="running",
                    worker_id=worker_id,
                    started_at=utcnow(),
                    finished_at=None,
                    error=None,
                    attempt_count=KnowledgeVectorSyncJob.attempt_count + 1,
                )
                .returning(
                    KnowledgeVectorSyncJob.version_id,
                    KnowledgeVectorSyncJob.operation,
                    KnowledgeVectorSyncJob.attempt_count,
                )
            )
            row = result.one_or_none()
            await db.commit()
        if row is None:
            return None
        return str(row.version_id), str(row.operation), int(row.attempt_count)

    async def _record_failure(
        self,
        job_id: str,
        version_id: str,
        attempt_count: int,
        worker_id: str,
        error: str,
    ) -> None:
        retrying = attempt_count < self._max_attempts
        async with self._session_factory() as db:
            job = await db.get(KnowledgeVectorSyncJob, job_id)
            if job is None or job.state != "running" or job.worker_id != worker_id:
                return
            job.state = "queued" if retrying else "failed"
            job.error = error
            job.worker_id = None
            job.finished_at = None if retrying else utcnow()
            version = await db.get(KnowledgeDocumentVersion, version_id)
            db.add(KnowledgeAuditEvent(
                event_type="vector_sync_failed",
                version_id=version_id,
                document_id=version.document_id if version is not None else None,
                actor_role="system",
                detail={"error": error[:500], "operation": job.operation, "will_retry": retrying},
            ))
            await db.commit()


def storage_key_for(version_id: str, filename: str) -> str:
    return f"documents/{version_id}/{safe_filename(filename)}"


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


knowledge_storage = KnowledgeStorage()
knowledge_vector_index = KnowledgeVectorIndex()
embedding_service = EmbeddingService()
knowledge_retrieval_service = KnowledgeRetrievalService(
    vector_index=knowledge_vector_index,
    embedding_service=embedding_service,
)
knowledge_ingestion_service = KnowledgeIngestionService(
    storage=knowledge_storage,
    vector_index=knowledge_vector_index,
    embedding_service=embedding_service,
)
knowledge_vector_sync_service = KnowledgeVectorSyncService(
    vector_index=knowledge_vector_index,
)
